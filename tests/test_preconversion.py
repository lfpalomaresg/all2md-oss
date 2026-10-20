# -*- coding: utf-8 -*-
"""Preconversión de formatos antiguos con LibreOffice (soffice): .xls →
.xlsx (para motor_excel, conserva combinadas) y .ppt → .pptx (para
Docling). Se salta si LibreOffice no está instalado. Datos sintéticos."""
import subprocess

import pytest
from openpyxl import Workbook

from preconversion import buscar_soffice, con_soffice

SOFFICE = buscar_soffice()
requiere_soffice = pytest.mark.skipif(SOFFICE is None, reason="LibreOffice no instalado")


def _a_formato(origen, formato, tmp_path):
    """Fixture: genera el .xls/.ppt de prueba. Con perfil propio, como
    con_soffice: con el perfil por defecto falla si LibreOffice está abierto."""
    perfil = (tmp_path / "perfil-fixture").resolve().as_uri()
    subprocess.run([SOFFICE, f"-env:UserInstallation={perfil}", "--headless",
                    "--convert-to", formato, "--outdir", str(tmp_path), str(origen)],
                   check=True, capture_output=True)
    return origen.with_suffix(f".{formato}")


def test_buscar_soffice_mira_rutas_conocidas_aunque_no_este_en_el_path(monkeypatch, tmp_path):
    # Lanzada desde Finder, la app no hereda el PATH de Homebrew.
    import preconversion
    falso = tmp_path / "soffice"
    falso.write_text("#!/bin/sh\n")
    falso.chmod(0o755)
    monkeypatch.setenv("PATH", "/usr/bin:/bin")
    monkeypatch.setattr(preconversion, "RUTAS_SOFFICE", [str(falso)])
    assert buscar_soffice() == str(falso)


def test_buscar_soffice_none_si_no_hay(monkeypatch):
    import preconversion
    monkeypatch.setenv("PATH", "/nonexistent")
    monkeypatch.setattr(preconversion, "RUTAS_SOFFICE", ["/nonexistent/soffice"])
    assert buscar_soffice() is None


@requiere_soffice
def test_xls_a_xlsx_conserva_celdas_combinadas(tmp_path):
    from motor_excel import convertir
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Concepto"
    ws["B1"] = "Mes ñ"
    ws.merge_cells("B1:C1")
    ws.append(["Rooms", 100, 120])
    fuente = tmp_path / "fuente.xlsx"
    wb.save(fuente)
    xls = _a_formato(fuente, "xls", tmp_path)
    assert xls.exists()

    xlsx = con_soffice(xls, "xlsx", tmp_path / "trabajo")
    md, _ = convertir(xlsx)
    assert "| Concepto | Mes ñ | Mes ñ |" in md
    assert "| Rooms | 100 | 120 |" in md


def test_con_soffice_origen_ilegible_da_error_legible(tmp_path, monkeypatch):
    # soffice, como textutil, importa la basura como texto en vez de fallar;
    # el fallo realista es no poder leer el origen.
    import preconversion
    monkeypatch.setattr(preconversion, "buscar_soffice", lambda: "/bin/true")
    with pytest.raises(RuntimeError, match="fantasma.xls"):
        con_soffice(tmp_path / "fantasma.xls", "xlsx", tmp_path / "trabajo")


@requiere_soffice
def test_motor_excel_cli_convierte_xls(tmp_path):
    import sys
    from pathlib import Path
    raiz = Path(__file__).resolve().parent.parent
    wb = Workbook()
    wb.active["A1"] = "Cifra sintética"
    wb.active["B1"] = 42
    wb.save(tmp_path / "f.xlsx")
    _a_formato(tmp_path / "f.xlsx", "xls", tmp_path)
    r = subprocess.run([sys.executable, str(raiz / "motor_excel.py"),
                        str(tmp_path / "f.xls"), str(tmp_path / "f.md")],
                       capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    assert "| Cifra sintética | 42 |" in (tmp_path / "f.md").read_text(encoding="utf-8")


@requiere_soffice
def test_preconvertir_ppt_a_pptx(tmp_path):
    from pptx import Presentation
    from motor_docling import preconvertir
    prs = Presentation()
    diapo = prs.slides.add_slide(prs.slide_layouts[1])
    diapo.shapes.title.text = "Diapositiva sintética"
    prs.save(tmp_path / "f.pptx")
    _a_formato(tmp_path / "f.pptx", "ppt", tmp_path)
    pptx = preconvertir(tmp_path / "f.ppt", tmp_path / "trabajo")
    assert pptx.suffix == ".pptx"
    assert Presentation(pptx).slides[0].shapes.title.text == "Diapositiva sintética"


@requiere_soffice
def test_xls_que_es_texto_no_filtra_el_nombre_temporal(tmp_path):
    # Muchos ".xls" de bancos son CSV/HTML: soffice los importa como texto y
    # llama a la hoja como el fichero. Debe ser el nombre del documento.
    from motor_excel import convertir
    falso = tmp_path / "extracto banco.xls"
    falso.write_text("Fecha;Importe\n2026-01-01;10\n", encoding="utf-8")
    md, _ = convertir(falso)
    assert "## entrada" not in md
    assert "## extracto banco" in md
