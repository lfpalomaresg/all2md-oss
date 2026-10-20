# -*- coding: utf-8 -*-
"""Tests de motor_excel con libros sintéticos generados en tmp_path."""
import datetime as dt
import re

import pytest

from motor_excel import _formatea


@pytest.mark.parametrize("valor,esperado", [
    (None, ""),
    (0.8769304604573576, "0.88"),
    (1234, "1234"),
    ("a|b", "a\\|b"),
    ("linea1\nlinea2", "linea1 linea2"),
    ("linea1\r\nlinea2", "linea1 linea2"),
    (dt.datetime(2026, 7, 1), "2026-07-01"),
    (dt.datetime(2026, 7, 1, 14, 30), "2026-07-01 14:30"),
    (dt.datetime(2026, 7, 1, 14, 30, 15), "2026-07-01 14:30:15"),
    (dt.date(2026, 7, 1), "2026-07-01"),
    (-0.001, "0.00"),
    (-0.0, "0.00"),
    (-12.5, "-12.50"),
])
def test_formatea(valor, esperado):
    assert _formatea(valor) == esperado


# --- Loop 4: convertir() de extremo a extremo con libros sintéticos ---
from openpyxl import Workbook

from motor_excel import ConversionVacia, convertir


def _guardar(wb, tmp_path, nombre="libro.xlsx"):
    ruta = tmp_path / nombre
    wb.save(ruta)
    return ruta


def test_convertir_propaga_cabeceras_combinadas(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = "Concepto"
    ws["B1"] = "Mes"
    ws.merge_cells("B1:C1")
    ws.append(["Habitaciones", 100, 120])
    md, avisos = convertir(_guardar(wb, tmp_path))
    assert "## Resumen" in md
    assert "| Concepto | Mes | Mes |" in md
    assert "| Habitaciones | 100 | 120 |" in md
    assert avisos == 0


def test_convertir_avisa_de_formulas_sin_valor_cacheado(tmp_path):
    # openpyxl no calcula fórmulas: al guardar no hay valor cacheado.
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["A2"] = "=A1*2"
    md, avisos = convertir(_guardar(wb, tmp_path))
    assert avisos == 1
    # La celda se marca: vacía, una IA podría leerla como un cero.
    assert "| (fórmula sin calcular) |" in md


def test_convertir_omite_hoja_por_encima_del_limite(tmp_path, monkeypatch):
    import motor_excel
    monkeypatch.setattr(motor_excel, "LIMITE_FILAS", 10)
    wb = Workbook()
    wb.active.title = "Normal"
    wb.active["A1"] = "ok"
    grande = wb.create_sheet("Gigante")
    grande["A11"] = "x"
    md, _ = convertir(_guardar(wb, tmp_path))
    assert "## Normal" in md
    assert "## Gigante" not in md
    assert "Gigante (11x1" in md


def test_convertir_libro_vacio_lanza_conversion_vacia(tmp_path):
    with pytest.raises(ConversionVacia):
        convertir(_guardar(Workbook(), tmp_path))


def test_convertir_descarta_filas_totalmente_vacias(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Bloque 1"
    ws["B1"] = 1
    ws["A4"] = "Bloque 2"
    ws["B4"] = 2
    md, _ = convertir(_guardar(wb, tmp_path))
    filas = [l for l in md.splitlines() if l.startswith("|")]
    assert filas == ["| Bloque 1 | 1 |", "| Bloque 2 | 2 |"]


def test_convertir_conserva_filas_con_ceros_falsos_y_formulas(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 0
    ws["A2"] = False
    ws["A3"] = "=1+1"  # sin valor cacheado: la fila no debe desaparecer
    md, _ = convertir(_guardar(wb, tmp_path))
    filas = [l for l in md.splitlines() if l.startswith("|")]
    assert filas == ["| 0 |", "| False |", "| (fórmula sin calcular) |"]


def _con_xml_parcheado(ruta, reemplazos):
    """Reescribe el XML de la hoja 1 -- simula un libro recalculado por
    Excel (openpyxl no calcula fórmulas, deja todos los <v> vacíos)."""
    import zipfile
    with zipfile.ZipFile(ruta) as z:
        contenido = {n: z.read(n) for n in z.namelist()}
    xml = contenido["xl/worksheets/sheet1.xml"].decode()
    # openpyxl serializa el <v> vacío como "<v></v>" o "<v />" según use
    # lxml o no: se normaliza para que el parche no dependa del entorno.
    xml = re.sub(r"<v\s*/>", "<v></v>", xml)
    for viejo, nuevo in reemplazos:
        assert viejo in xml
        xml = xml.replace(viejo, nuevo)
    contenido["xl/worksheets/sheet1.xml"] = xml.encode()
    with zipfile.ZipFile(ruta, "w") as z:
        for n, datos in contenido.items():
            z.writestr(n, datos)
    return ruta


def test_formula_calculada_que_devuelve_cadena_vacia_no_se_marca(tmp_path):
    # =IF(...;"";x) recalculado devuelve "" -- data_only lo lee como None,
    # igual que una fórmula sin calcular. Si el libro SÍ se recalculó
    # (alguna fórmula trae valor), ese None es un "" legítimo.
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Total"
    ws["B1"] = '=IF(1,"",1)'
    ws["C1"] = "=1+1"
    ruta = _con_xml_parcheado(_guardar(wb, tmp_path), [
        ('<c r="B1"><f>', '<c r="B1" t="str"><f>'),
        ('<f>1+1</f><v></v>', '<f>1+1</f><v>2</v>'),
    ])
    md, avisos = convertir(ruta)
    # La columna del "" queda vacía y se descarta; lo importante: sin marca.
    assert "| Total | 2 |" in md
    assert "sin calcular" not in md
    assert avisos == 0


def test_formula_sin_calcular_en_cabecera_combinada_marca_todas_sus_celdas(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "=1+1"
    ws.merge_cells("A1:B1")
    ws["A2"] = "x"
    ws["B2"] = "y"
    md, _ = convertir(_guardar(wb, tmp_path))
    assert "| (fórmula sin calcular) | (fórmula sin calcular) |" in md


def test_formula_matricial_sin_calcular_se_detecta(tmp_path):
    # Con data_only=False openpyxl devuelve un objeto ArrayFormula, no un
    # str "=..." -- antes no se marcaba ni se contaba.
    from openpyxl.worksheet.formula import ArrayFormula
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["B1"] = ArrayFormula("B1", "=SUM(A1:A1*2)")
    md, avisos = convertir(_guardar(wb, tmp_path))
    assert avisos == 1
    assert "| 1 | (fórmula sin calcular) |" in md


# --- Loop 9: columnas espaciadoras vacías ---
def test_convertir_descarta_columnas_totalmente_vacias(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Concepto"
    ws["C1"] = "Real"      # B vacía: columna espaciadora
    ws["E1"] = "Budget"    # D vacía
    ws["A2"] = "Rooms"
    ws["C2"] = 10
    ws["E2"] = 0
    md, _ = convertir(_guardar(wb, tmp_path))
    filas = [l for l in md.splitlines() if l.startswith("|")]
    assert filas == ["| Concepto | Real | Budget |", "| Rooms | 10 | 0 |"]


def test_columna_cubierta_por_combinada_no_se_descarta(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Mes"
    ws.merge_cells("A1:B1")
    ws["A2"] = 1          # B2 vacía, pero B1 hereda "Mes"
    md, _ = convertir(_guardar(wb, tmp_path))
    assert "| Mes | Mes |" in md


def test_combinada_enorme_no_se_salta_el_techo(tmp_path):
    # max_column de openpyxl ya cuenta las celdas combinadas: la hoja se
    # omite por el techo antes de que la ampliación de rango la vea.
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Título"
    ws.merge_cells("A1:XFD1")
    ws["A2"] = 1
    md, _ = convertir(_guardar(wb, tmp_path))
    assert "(2x16384, supera el límite" in md
    assert "| Título" not in md


# --- Loop 10: hojas ocultas señaladas ---
@pytest.mark.parametrize("estado,encabezado", [
    ("hidden", "## Aux (hoja oculta)"),
    ("veryHidden", "## Aux (hoja oculta)"),
    ("visible", "## Aux\n"),
])
def test_hoja_oculta_se_senala_en_el_encabezado(tmp_path, estado, encabezado):
    wb = Workbook()
    wb.active["A1"] = "visible"
    aux = wb.create_sheet("Aux")
    aux["A1"] = "tabla auxiliar"
    aux.sheet_state = estado
    md, _ = convertir(_guardar(wb, tmp_path))
    assert encabezado in md


# --- Automejora 2, loop 2: barra invertida justo antes de un pipe ---
@pytest.mark.parametrize("valor", ["a\\|b", "a\\\\|b", "a\\", "C:\\ruta"])
def test_barra_invertida_no_rompe_la_rejilla(valor):
    # Tokenizador GFM: "\x" es un escape (incluido "\|"); un "|" suelto
    # separa celdas. Un "\|" literal salía como "\\|" = barra escapada +
    # pipe separador, y la fila ganaba una columna.
    fila = "| " + _formatea(valor) + " | fin |"
    celdas = [c.strip() for c in re.findall(r"(?:\\.|[^\\|])+", fila)]
    celdas = [c for c in celdas if c]
    assert len(celdas) == 2 and celdas[1] == "fin", fila


# --- Automejora 2, loop 3: todos los saltos de línea, no solo \n y \r ---
@pytest.mark.parametrize("sep", ["\x0b", "\x0c", "\x1c", "\x1d", "\x1e", "\x85", " ", " "])
def test_otros_saltos_de_linea_no_parten_la_fila(sep):
    # str.splitlines() (y muchos lectores de Markdown/RAG) cortan la línea
    # en estos caracteres: una celda con ellos partía la fila en dos.
    assert _formatea(f"a{sep}b") == "a b"


# --- Automejora 2, loop 4: celdas de solo espacios ---
def test_celdas_de_solo_espacios_cuentan_como_vacias(tmp_path):
    # En plantillas reales se "borra" una celda tecleando un espacio: esa
    # fila/columna espaciadora no debe sobrevivir al filtro de vacías.
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Concepto"
    ws["B1"] = "   "        # columna B: solo espacios
    ws["C1"] = "Real"
    ws["A2"] = " "          # fila 2: solo espacios
    ws["A3"] = "Rooms"
    ws["C3"] = 10
    md, _ = convertir(_guardar(wb, tmp_path))
    filas = [l for l in md.splitlines() if l.startswith("|")]
    assert filas == ["| Concepto | Real |", "| Rooms | 10 |"]


def test_la_sangria_de_una_etiqueta_se_conserva():
    # En un USALI la sangría marca jerarquía (subpartidas): no se recorta.
    assert _formatea("   Rooms Revenue") == "   Rooms Revenue"


# --- Automejora 2, loop 5: hoja que queda vacía tras filtrar ---
def test_hoja_de_solo_espacios_no_deja_encabezado_huerfano(tmp_path):
    wb = Workbook()
    wb.active.title = "Datos"
    wb.active["A1"] = "ok"
    blanca = wb.create_sheet("Blanca")
    blanca["A1"] = "   "
    md, _ = convertir(_guardar(wb, tmp_path))
    assert "## Datos" in md
    assert "Blanca" not in md


def test_libro_de_solo_espacios_es_conversion_vacia(tmp_path):
    wb = Workbook()
    wb.active["B2"] = " "
    with pytest.raises(ConversionVacia):
        convertir(_guardar(wb, tmp_path))


# --- Automejora 2, loop 6: horas y duraciones ---
@pytest.mark.parametrize("valor,esperado", [
    (dt.time(14, 30), "14:30"),
    (dt.time(14, 30, 15), "14:30:15"),
    (dt.time(0, 0), "00:00"),
    (dt.timedelta(hours=26), "26:00"),          # formato [h]:mm de Excel
    (dt.timedelta(hours=1, minutes=5, seconds=9), "1:05:09"),
    (dt.timedelta(minutes=-90), "-1:30"),
])
def test_formatea_horas_y_duraciones(valor, esperado):
    # str() daba "14:30:00" y "1 day, 2:00:00": ruido y, en la duración,
    # inglés mezclado en el .md.
    assert _formatea(valor) == esperado


# --- Automejora 2, loop 7: celdas con formato de porcentaje ---
def test_porcentaje_sale_como_porcentaje(tmp_path):
    # Ocupación 0.8769 con formato "0.0%": Excel enseña 87.7%, y un "0.88"
    # suelto se puede leer como 0,88 habitaciones o euros.
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Ocupación"
    ws["B1"] = 0.8769304604573576
    ws["B1"].number_format = "0.0%"
    ws["C1"] = 1
    ws["C1"].number_format = "0%"
    ws["D1"] = 0.5          # mismo número sin formato %: no se toca
    md, _ = convertir(_guardar(wb, tmp_path))
    assert "| Ocupación | 87.69% | 100.00% | 0.50 |" in md


def test_porcentaje_en_combinada_se_propaga_con_su_formato(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 0.25
    ws["A1"].number_format = "0%"
    ws.merge_cells("A1:B1")
    md, _ = convertir(_guardar(wb, tmp_path))
    assert "| 25.00% | 25.00% |" in md


@pytest.mark.parametrize("formato", ['0.00" %"', '"%"0', "General", "@"])
def test_un_porcentaje_literal_entre_comillas_no_multiplica(formato):
    # Un "%" entre comillas es texto decorativo, no formato de porcentaje.
    assert _formatea(0.5, formato) == "0.50"


# --- Automejora 2, loop 16: plantillas .xltx ---
def test_convertir_lee_una_plantilla_xltx(tmp_path):
    wb = Workbook()
    wb.active["A1"] = "Plantilla"
    wb.template = True
    ruta = tmp_path / "cierre.xltx"
    wb.save(ruta)
    md, _ = convertir(ruta)
    assert "| Plantilla |" in md


# --- Ronda 3, loop E: formatos de número, horas y celdas invisibles ---
@pytest.mark.parametrize("valor,formato,esperado", [
    (0.5, "[$%-409]0.00", "0.50"),          # % dentro de una etiqueta [..]
    (0.5, "0.00;-0.00;0.00;@%", "0.50"),    # % solo en la sección de texto
    (0.5, "0.00;-0.00;0%", "0.50"),         # % solo en la sección del cero
    (0, "0.00;-0.00;0%", "0.00%"),
    (-0.05, "0.00;-0.0%", "-5.00%"),        # % solo en la sección negativa
    (0.05, "0.00;-0.0%", "0.05"),
    (0.25, "0%", "25.00%"),
])
def test_porcentaje_segun_la_seccion_que_aplica(valor, formato, esperado):
    assert _formatea(valor, formato) == esperado


@pytest.mark.parametrize("valor,formato,esperado", [
    (dt.time(23, 59, 59, 999000), "h:mm", "00:00"),     # Excel redondea al segundo
    (dt.time(23, 59, 59, 999000), "h:mm:ss", "00:00:00"),
    (dt.time(10, 29, 45), "h:mm", "10:29"),             # el formato no lleva segundos
    (dt.time(10, 29, 45), "General", "10:29:45"),
    (dt.datetime(1900, 1, 1, 6, 0), "h:mm", "06:00"),   # 1.25 con h:mm: Excel enseña 6:00
    (dt.datetime(1900, 1, 1, 6, 0), "dd/mm/yyyy h:mm", "1900-01-01 06:00"),
    (dt.time(0, 0), "dd/mm/yyyy", "0"),                 # un 0 en columna de fechas
    (dt.time(0, 0), "h:mm", "00:00"),
])
def test_horas_segun_formato(valor, formato, esperado):
    assert _formatea(valor, formato) == esperado


@pytest.mark.parametrize("valor", ["​", "﻿", " ​ "])
def test_celda_solo_invisible_cuenta_como_vacia(valor):
    assert _formatea(valor) == ""


@pytest.mark.parametrize("formato,es_fecha", [
    ("General", False), ("@", False), ("0.00", False), ("h:mm", False),
    ("dd/mm/yyyy", True), ("yyyy-mm-dd h:mm", True), ('0.00" días"', False),
])
def test_es_formato_de_fecha(formato, es_fecha):
    from motor_excel import _es_formato_de_fecha
    assert _es_formato_de_fecha(formato) is es_fecha


@pytest.mark.parametrize("formato", ["mmmm", "mmm", "MMMM", "m", "mmmm;@", "[$-es-ES]mmmm", "mmm-yy"])
def test_fecha_con_formato_de_mes_no_se_toma_por_hora(formato):
    # Cabeceras de mes de un USALI: openpyxl da datetime(2024, 7, 1).
    # Tomarlas por "hora sin fecha" las convertía todas en "00:00" (careo, ALTA).
    assert _formatea(dt.datetime(2024, 7, 1), formato) == "2024-07-01"


@pytest.mark.parametrize("valor,formato,esperado", [
    (dt.time(0, 1, 30), "mm:ss", "01:30"),          # sin horas en el formato
    (dt.time(0, 0, 0, 500000), "h:mm:ss", "00:00:01"),  # medio segundo hacia arriba, como Excel
    (dt.datetime(2024, 7, 1, 6, 0), "dd/mm/yyyy hh:mm", "2024-07-01 06:00"),
])
def test_horas_casos_finos(valor, formato, esperado):
    assert _formatea(valor, formato) == esperado


def test_porcentaje_en_formato_condicional_elige_la_seccion_por_condicion():
    # [>=100]0;0.0% con 50: Excel aplica la 2ª sección (por condición, no
    # por signo) y enseña 5000.0%.
    assert _formatea(0.5, "[>=100]0;0.0%") == "50.00%"
    assert _formatea(50, "[>=100]0;0.0%") == "5000.00%"


# --- Pendientes ronda 3: condiciones, AM/PM, décimas, matriciales ---
@pytest.mark.parametrize("valor,formato,esperado", [
    (5, "[<1]0.0%;0", "5"),                   # no cumple [<1] → 2ª sección, sin %
    (0.5, "[<1]0.0%;0", "50.00%"),
    (50, "[>=100]0;0.0%", "5000.00%"),
    (150, "[>=100]0;0.0%", "150"),
    (150.0, '[>100]0.0%;[<0]0;0', "15000.00%"),
    (-3, '[>100]0.0%;[<0]0;0.0%', "-3"),
    (0.5, '[>100]0;[<0]0;0.0%', "50.00%"),    # ninguna condición → 3ª sección
    (0.5, '[Red][<1]0.0%;0', "50.00%"),       # color + condición
])
def test_porcentaje_en_formato_condicional_evalua_la_condicion(valor, formato, esperado):
    assert _formatea(valor, formato) == esperado


@pytest.mark.parametrize("valor,formato,esperado", [
    (dt.time(14, 24), "hh:mm AM/PM", "02:24 PM"),
    (dt.time(0, 5), "h:mm AM/PM", "12:05 AM"),
    (dt.time(12, 0), "h:mm AM/PM", "12:00 PM"),
    (dt.time(14, 24, 5), "h:mm:ss AM/PM", "02:24:05 PM"),
    (dt.time(14, 24), "h:mm A/P", "02:24 P"),
    (dt.time(0, 1, 30, 500000), "mm:ss.0", "01:30.5"),
    (dt.time(0, 1, 30, 456000), "mm:ss.00", "01:30.46"),
    (dt.time(0, 1, 59, 960000), "mm:ss.0", "02:00.0"),   # el redondeo arrastra
    (dt.time(10, 0, 0, 250000), "hh:mm:ss.000", "10:00:00.250"),
])
def test_horas_am_pm_y_fracciones(valor, formato, esperado):
    assert _formatea(valor, formato) == esperado


@pytest.mark.parametrize("valor,formato,esperado", [
    (-0.5, "0.00\\;0;0.0%", "-50.00%"),   # "\\;" es un ";" literal, no separa
    (0.5, "0\\;0.0%", "50.00%"),
])
def test_punto_y_coma_escapado_no_separa_secciones(valor, formato, esperado):
    assert _formatea(valor, formato) == esperado


@pytest.mark.parametrize("valor,formato,esperado", [
    (dt.time(18, 30, 15), "ss.00", "15.00"),       # solo segundos
    (dt.time(14, 0), "h AM/PM", "2 PM"),           # solo la hora
    (dt.time(9, 0), "hh AM/PM", "09 AM"),
    (dt.timedelta(hours=30, microseconds=450000), "[h]:mm:ss.0", "30:00:00.5"),
    (dt.timedelta(minutes=15, seconds=2), "[mm]:ss", "15:02"),
    (dt.timedelta(minutes=75, seconds=2, microseconds=456000), "[mm]:ss.00", "75:02.46"),
    (dt.timedelta(seconds=3725), "[ss]", "3725"),
    (dt.timedelta(hours=26), "[h]:mm", "26:00"),
    (dt.timedelta(minutes=-90), "[h]:mm", "-1:30"),
])
def test_horas_y_duraciones_con_unidades_parciales(valor, formato, esperado):
    assert _formatea(valor, formato) == esperado


@pytest.mark.parametrize("valor,formato,esperado", [
    (dt.timedelta(seconds=5415, microseconds=260000), "[ss].00", "5415.26"),
    (dt.timedelta(minutes=5, seconds=3), "[m]:ss", "5:03"),
    (dt.timedelta(hours=30), "[h]", "30"),
])
def test_duraciones_formatos_acumulados_finos(valor, formato, esperado):
    assert _formatea(valor, formato) == esperado


def test_formula_matricial_multicelda_sin_calcular_marca_todo_su_rango(tmp_path):
    # openpyxl solo pone el ArrayFormula en la celda ancla; las demás del
    # rango llegan vacías y salían en blanco (pendiente de la 1ª ronda).
    from openpyxl.worksheet.formula import ArrayFormula
    wb = Workbook()
    ws = wb.active
    for fila, v in enumerate([1, 2, 3], start=1):
        ws.cell(row=fila, column=1, value=v)
    ws["B1"] = ArrayFormula("B1:B3", "=A1:A3*2")
    md, avisos = convertir(_guardar(wb, tmp_path))
    filas = [l for l in md.splitlines() if l.startswith("|")]
    assert filas == [f"| {v} | (fórmula sin calcular) |" for v in (1, 2, 3)]
    assert avisos == 3


def test_formula_matricial_con_rango_enorme_no_cuelga(tmp_path):
    import time
    from openpyxl.worksheet.formula import ArrayFormula
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["B1"] = ArrayFormula("B1:XFD1048576", "=A1*2")
    inicio = time.monotonic()
    md, _ = convertir(_guardar(wb, tmp_path))
    import motor_excel
    assert time.monotonic() - inicio < 10
    assert md.count("(fórmula sin calcular)") <= motor_excel.MAX_CELDAS_MATRICIALES + 1


def _libro_matricial(celdas, matriciales, combinar=()):
    from openpyxl.worksheet.formula import ArrayFormula
    wb = Workbook()
    ws = wb.active
    for ref, v in celdas.items():
        ws[ref] = v
    for ancla, (rango, formula) in matriciales.items():
        ws[ancla] = ArrayFormula(rango, formula)
    for rango in combinar:
        ws.merge_cells(rango)
    return wb


def test_matricial_solapada_con_combinada_no_rompe_la_conversion(tmp_path):
    wb = _libro_matricial({"A1": 1, "A2": 2, "A3": 3}, {"B1": ("B1:B3", "=A1:A3*2")},
                          combinar=["A2:B2"])
    md, _ = convertir(_guardar(wb, tmp_path))   # antes: AttributeError (MergedCell)
    assert "(fórmula sin calcular)" in md


def test_matricial_mas_alla_de_los_datos_se_marca_entera(tmp_path):
    wb = _libro_matricial({"A1": 1, "A2": 2, "A3": 3}, {"B1": ("B1:D1", "=TRANSPOSE(A1:A3)")})
    md, avisos = convertir(_guardar(wb, tmp_path))
    marca = "(fórmula sin calcular)"
    assert f"| 1 | {marca} | {marca} | {marca} |" in md
    assert avisos == 3


def test_matricial_sin_duplicados_en_el_recuento(tmp_path):
    wb = _libro_matricial({"A1": 1, "A2": 2, "A3": 3, "B2": "=A2"},
                          {"B1": ("B1:B3", "=A1:A3*2")})
    _, avisos = convertir(_guardar(wb, tmp_path))
    assert avisos == 3


def test_matricial_cuyo_rango_no_empieza_en_la_ancla_no_se_expande(tmp_path):
    wb = _libro_matricial({"A1": 1, "A2": 2, "A3": 3}, {"B2": ("B1:B3", "=A1:A3*2")})
    md, avisos = convertir(_guardar(wb, tmp_path))
    assert avisos == 1
    assert md.count("(fórmula sin calcular)") == 1


def test_muchas_matriciales_hostiles_no_cuelgan(tmp_path):
    import time
    import motor_excel
    celdas = {f"A{f}": f for f in range(1, 201)}
    matriciales = {f"B{f}": (f"B{f}:XFD1048576", "=1") for f in range(1, 201)}
    wb = _libro_matricial(celdas, matriciales)
    inicio = time.monotonic()
    _, avisos = convertir(_guardar(wb, tmp_path))
    assert time.monotonic() - inicio < 15
    assert avisos <= motor_excel.MAX_CELDAS_MATRICIALES + 200


def test_matriciales_hostiles_sobre_celdas_con_datos_no_multiplican_el_trabajo(tmp_path):
    # Celdas ya ocupadas no se marcan, pero recorrerlas también cuesta.
    import time
    celdas = {f"{col}{f}": 1 for f in range(1, 101) for col in ("A", "B", "C", "D", "E")}
    matriciales = {f"F{f}": (f"F{f}:SF5000", "=1") for f in range(1, 101)}
    for f in range(1, 101):
        for col in ("G", "H", "I"):
            celdas[f"{col}{f}"] = 1
    wb = _libro_matricial(celdas, matriciales)
    inicio = time.monotonic()
    convertir(_guardar(wb, tmp_path))
    assert time.monotonic() - inicio < 5


def test_matricial_mayor_que_el_presupuesto_marca_solo_la_ancla_y_avisa(tmp_path, monkeypatch):
    import motor_excel
    monkeypatch.setattr(motor_excel, "MAX_CELDAS_MATRICIALES", 5)
    wb = _libro_matricial({"A1": 1}, {"B1": ("B1:B10", "=1")})
    md, avisos = convertir(_guardar(wb, tmp_path))
    assert md.count("(fórmula sin calcular)") == 1 and avisos == 1
    assert "B1:B10" in md and "AVISO" in md
