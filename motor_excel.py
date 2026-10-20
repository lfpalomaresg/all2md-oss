"""
Conversor de Excel nativo (.xlsx/.xlsm con celdas combinadas) a Markdown SIN
IA -- lee directamente las celdas combinadas (merged_cells) de openpyxl y
propaga el valor de cada cabecera combinada a todas las columnas que cubre.
Un Excel nativo trae esa información exacta; no hace falta que nadie la
"adivine" a partir de una imagen.

Motor de Excel utilizado por all2md_repo_app.py. Incluye tratamiento de
celdas combinadas y comprobaciones de formatos, fórmulas y límites.

No produce una tabla Markdown de cabecera única: un USALI (y en general
un cierre financiero real) tiene varias filas de cabecera apiladas, así
que se vuelca la rejilla completa con sintaxis de pipes para legibilidad,
sin la fila separadora `|---|` de una tabla de cabecera única (no
aplicaría: no hay UNA fila de cabecera).

CLI: python motor_excel.py <origen.xlsx> <destino.md>
Librería: from motor_excel import convertir
"""
import datetime as dt
import re
import sys
import tempfile
import unicodedata
from pathlib import Path

from all2md_core import escribir_atomico, misma_ruta
from preconversion import con_soffice

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import range_boundaries
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

# Techo de seguridad ANTES de escanear ninguna celda: openpyxl reporta
# max_row/max_column hasta la última celda con formato aplicado, no la
# última con contenido -- un XLSX con formato pintado hasta el límite de
# Excel (1048576 x 16384) convertiría el escaneo en decenas de millones de
# celdas y colgaría el proceso (hallazgo de careo, ALTA, 23/08: el primer
# intento de acotar el rango seguía iterando sin límite ANTES de saber
# dónde parar). Un USALI real ronda 100 filas x 100 columnas por hoja --
# esto deja margen generoso sin abrir la puerta a un XLSX hostil.
LIMITE_FILAS = 5000
LIMITE_COLUMNAS = 500

# Una fórmula sin valor cacheado (libro guardado sin recalcular) sale así
# en vez de vacía: una celda vacía en un cierre financiero puede leerse
# como un cero, y esto no lo es.
MARCA_FORMULA = "(fórmula sin calcular)"

# Celdas que se pueden recorrer por hoja al expandir matriciales
# multicelda: una real rara vez pasa de unos cientos; esto acota el trabajo
# (y las marcas) de un libro hostil.
MAX_CELDAS_MATRICIALES = 10_000


class ConversionVacia(Exception):
    """El libro no produjo ningún contenido."""


class HojaDemasiadoGrande(Exception):
    """Una hoja supera el techo de seguridad -- no se procesa esa hoja."""


def _limpia_formato(formato: str) -> str:
    """Formato numérico sin lo que es texto literal: lo entrecomillado
    (0.00" %"), lo escapado con barra (0.00\\%) y las etiquetas entre
    corchetes ([$%-409], [Red], [h])."""
    return re.sub(r'"[^"]*"|\\.|\[[^\]]*\]', "", formato or "")


_CONDICION = re.compile(r"\[(<=|>=|<>|<|>|=)\s*(-?\d+(?:\.\d+)?)\]")
_OPERADORES = {
    "<": lambda a, b: a < b, ">": lambda a, b: a > b, "=": lambda a, b: a == b,
    "<=": lambda a, b: a <= b, ">=": lambda a, b: a >= b, "<>": lambda a, b: a != b,
}


def _secciones(formato: str) -> list[str]:
    # ";" separa secciones salvo dentro de comillas ("a;b") o escapado (\;).
    secciones, actual = [], ""
    for trozo in re.findall(r'"[^"]*"?|\\.?|;|[^";\\]+', formato or ""):
        if trozo == ";":
            secciones.append(actual)
            actual = ""
        else:
            actual += trozo
    return secciones + [actual]


def _cumple(seccion: str, v) -> bool | None:
    """True/False si la sección lleva condición ([>=100]...), None si no."""
    m = _CONDICION.search(seccion)
    if not m:
        return None
    return _OPERADORES[m.group(1)](v, float(m.group(2)))


def _seccion(formato: str, v) -> str:
    """La sección del formato que Excel aplica a `v` (ya sin literales).
    Sin condiciones, por signo: "pos;neg;cero;texto" -- un "%" en la
    sección del cero o del texto no convierte un 0,5 en 50% (hallazgo de
    careo, BAJA, 04/10). Con condiciones ([<1]0.0%;0), como Excel: la 1ª
    si cumple la suya, si no la 2ª (si cumple o no tiene), si no la 3ª."""
    secciones = _secciones(formato)
    condiciones = [_cumple(s, v) for s in secciones[:2]]
    if any(c is not None for c in condiciones):
        if condiciones[0]:
            elegida = secciones[0]
        elif len(secciones) > 1 and condiciones[1] is not False:
            elegida = secciones[1]
        elif len(secciones) > 2:
            elegida = secciones[2]
        else:
            elegida = secciones[0]
        return _limpia_formato(elegida)
    secciones = [_limpia_formato(s) for s in secciones]
    if v < 0 and len(secciones) > 1:
        return secciones[1]
    if v == 0 and len(secciones) > 2:
        return secciones[2]
    return secciones[0]


def _es_porcentaje(formato: str, v=1) -> bool:
    return "%" in _seccion(formato, v)


def _tokens(formato: str) -> str:
    limpio = _limpia_formato(formato).lower()
    return "" if limpio in ("", "general", "@") else limpio


def _es_solo_hora(formato: str) -> bool:
    """h:mm, mm:ss... con horas o segundos y sin día ni año. "m" a secas es
    MES: "mmmm" es una cabecera de mes, no una hora (hallazgo de careo,
    ALTA, 04/10: salía "00:00")."""
    t = _tokens(formato)
    return bool(re.search(r"[hs]", t)) and not re.search(r"[dye]", t)


def _es_formato_de_fecha(formato: str) -> bool:
    # Día, año ("e": año de era) o mes. "General" lleva una "e" y no es fecha.
    t = _tokens(formato)
    return bool(re.search(r"[dyem]", t)) and not _es_solo_hora(formato)


def _hora(v: dt.time, formato: str) -> str:
    """Hora como la enseña Excel con su formato. Redondeo al segundo (o a
    la fracción que muestre ss.0/ss.00), medio hacia arriba: openpyxl deja
    0.99999999 en 23:59:59.999, que Excel enseña como 0:00 (round() de
    Python va al par). Segundos solo si el formato los muestra o, en
    General, si los hay; mm:ss sin horas; AM/PM y A/P en 12 h."""
    t = _tokens(formato)
    general = not t
    m = re.search(r"s\.(0+)", t)
    decimales = len(m.group(1)) if m else 0
    unidades = 10 ** decimales
    exacto = (v.hour * 3600 + v.minute * 60 + v.second) * unidades + v.microsecond * unidades / 1e6
    total = int(exacto + 0.5) % (86400 * unidades)
    total, fraccion = divmod(total, unidades)
    horas, resto = divmod(total, 3600)
    minutos, segundos = divmod(resto, 60)

    sufijo = ""
    doce_horas = "am/pm" in t or "a/p" in t
    if doce_horas:
        tarde = horas >= 12
        horas = horas % 12 or 12
        sufijo = (" PM" if tarde else " AM") if "am/pm" in t else (" P" if tarde else " A")

    seg = f"{segundos:02d}" + (f".{fraccion:0{decimales}d}" if decimales else "")
    # La "m" de "am/pm" no son minutos.
    t = t.replace("am/pm", "").replace("a/p", "")
    if not general and "h" not in t and "m" not in t and "s" in t:
        return seg  # ss.00: solo segundos
    if not general and "h" in t and "m" not in t and "s" not in t:
        return f"{horas:02d}{sufijo}" if "hh" in t else f"{horas}{sufijo}"  # h AM/PM
    if not general and "h" not in t and "s" in t:
        return f"{minutos:02d}:{seg}"  # mm:ss
    if "s" in t or (general and segundos):
        return f"{horas:02d}:{minutos:02d}:{seg}{sufijo}"
    return f"{horas:02d}:{minutos:02d}{sufijo}"


def _duracion(v: dt.timedelta, formato: str) -> str:
    """Duración como los formatos acumulados de Excel: [h]:mm (por defecto,
    sin el "1 day, 2:00:00" de str()), [mm]:ss en minutos totales, [ss] en
    segundos totales, con las fracciones de ss.0 redondeadas hacia arriba."""
    t = _tokens(formato)
    bruto = (formato or "").lower()
    m = re.search(r"s\]?\.(0+)", bruto)  # ss.0 o [ss].00
    decimales = len(m.group(1)) if m else 0
    unidades = 10 ** decimales
    exacto = v.total_seconds() * unidades
    total = int(abs(exacto) + 0.5)
    signo = "-" if exacto < 0 and total else ""
    total, fraccion = divmod(total, unidades)
    frac = f".{fraccion:0{decimales}d}" if decimales else ""
    if re.search(r"\[s+\]", bruto):
        return f"{signo}{total}{frac}"
    minutos_totales, segundos = divmod(total, 60)
    acumulado = re.search(r"\[(m+)\]", bruto)
    if acumulado:
        relleno = "02d" if len(acumulado.group(1)) > 1 else "d"
        return f"{signo}{minutos_totales:{relleno}}:{segundos:02d}{frac}"
    horas, minutos = divmod(minutos_totales, 60)
    if re.fullmatch(r"(\[\$[^\]]*\])?\[h+\]", bruto.strip()):  # [h] a secas
        return f"{signo}{horas}"
    if "s" in t or (not t and (segundos or fraccion)):
        return f"{signo}{horas}:{minutos:02d}:{segundos:02d}{frac}"
    return f"{signo}{horas}:{minutos:02d}"


def _formatea(v, formato="General"):
    if v is None:
        return ""
    if (isinstance(v, (int, float)) and not isinstance(v, bool)
            and _es_porcentaje(formato, v)):
        # Ocupación 0.8769 con formato "0.0%": Excel enseña 87.7%. Un
        # "0.88" suelto se podía leer como 0,88 unidades.
        s = f"{v * 100:.2f}%"
        if s == "-0.00%":
            s = "0.00%"
    elif isinstance(v, float):
        # Redondeo a 2 decimales: la precisión binaria completa
        # (0.8769304604573576) no aporta nada a una IA/RAG y triplica
        # el tamaño del fichero sin motivo.
        s = f"{v:.2f}"
        if s == "-0.00":  # -0.001 redondeado: el signo solo confunde
            s = "0.00"
    elif isinstance(v, dt.datetime) and _es_solo_hora(formato):
        # Hora sin fecha (h:mm) con un valor de 24 h o más: openpyxl lo da
        # como datetime de 1900; Excel enseña solo la hora.
        s = _hora(v.time(), formato)
    elif isinstance(v, dt.time) and v == dt.time(0, 0) and _es_formato_de_fecha(formato):
        # Un 0 en una columna con formato de fecha: openpyxl da time(0, 0),
        # y "00:00" se leería como una hora. Es un cero.
        s = "0"
    elif isinstance(v, dt.datetime):
        # openpyxl devuelve datetime para cualquier celda con formato de
        # fecha: "2026-07-01 00:00:00" es ruido -- la hora solo si la hay.
        if v.time() == dt.time(0, 0):
            s = v.strftime("%Y-%m-%d")
        elif v.second or v.microsecond:
            s = v.strftime("%Y-%m-%d %H:%M:%S")
        else:
            s = v.strftime("%Y-%m-%d %H:%M")
    elif isinstance(v, dt.time):
        s = _hora(v, formato)
    elif isinstance(v, dt.timedelta):
        s = _duracion(v, formato)
    else:
        s = str(v)
    # Escapar lo que rompería la rejilla de pipes: un "|" o un salto de
    # línea dentro de una celda desplazaría las columnas de toda esa fila.
    # Las barras justo antes de un "|" se doblan: un "\|" literal salía
    # como "\\|" (barra escapada + pipe separador). El resto de barras
    # se deja tal cual (C:\ruta) para no ensuciar el texto del RAG.
    s = re.sub(r"(\\*)\|", lambda m: m.group(1) * 2 + "\\|", s)
    # splitlines() corta en TODOS los saltos que reconoce Python (\v, \f,
    # \x1c-\x1e, \x85, \u2028, \u2029, además de \n y \r): un lector que
    # parta por líneas no debe ver la fila de una celda partida en dos.
    s = " ".join(s.splitlines())
    # Una celda de solo espacios es una celda "borrada" a mano: cuenta como
    # vacía para el filtro de filas/columnas espaciadoras. La sangría de
    # una etiqueta con texto se conserva (marca jerarquía en un USALI).
    # Invisibles (ZWSP, BOM) también: una celda con solo eso mantenía viva
    # una fila o columna separadora (hallazgo de careo, BAJA, 04/10).
    visible = any(not ch.isspace() and unicodedata.category(ch) != "Cf" for ch in s)
    return s if visible else ""


def _escanear(hoja, hoja_formulas):
    """Recorre la hoja UNA sola vez y devuelve (max_fila, max_col,
    celdas_formula_sin_valor, n_formulas_con_valor, refs de matriciales
    demasiado grandes para expandir). Una celda cuenta como
    "contenido real" si tiene valor en la copia data_only=True, O si es una
    fórmula sin valor cacheado (data_only=True la ve como None, pero
    data_only=False ve el texto "=..."). Sin esto último, una hoja hecha
    solo de fórmulas sin recalcular se recortaba a un rango vacío y el aviso
    nunca llegaba a emitirse (hallazgo de careo, MEDIA, 23/08)."""
    max_fila = max_col = 0
    sin_valor = {}  # (fila, col) -> None: conjunto ordenado, sin duplicados
    con_valor = 0
    presupuesto = MAX_CELDAS_MATRICIALES
    grandes = []
    for fila_datos, fila_formulas in zip(hoja.iter_rows(), hoja_formulas.iter_rows()):
        for c, cf in zip(fila_datos, fila_formulas):
            # Las matriciales y las tablas de datos no llegan como str
            # "=...", sino como objeto -- sin esto no se marcaban ni contaban.
            es_formula = (
                isinstance(cf.value, (ArrayFormula, DataTableFormula))
                or (isinstance(cf.value, str) and cf.value.startswith("="))
            )
            if c.value is not None or es_formula:
                max_fila = max(max_fila, c.row)
                max_col = max(max_col, c.column)
            if not es_formula:
                continue
            if c.value is not None:
                con_valor += 1
                continue
            sin_valor[(c.row, c.column)] = None
            if isinstance(cf.value, ArrayFormula) and cf.value.ref:
                presupuesto, f_max, c_max = _expandir_matricial(
                    hoja, c, cf.value.ref, sin_valor, presupuesto, grandes)
                max_fila, max_col = max(max_fila, f_max), max(max_col, c_max)
    return max_fila, max_col, list(sin_valor), con_valor, grandes


def _expandir_matricial(hoja, ancla, ref, sin_valor, presupuesto, grandes):
    """Una matricial multicelda solo lleva el ArrayFormula en la ancla: el
    resto de su rango llega vacío y salía en blanco en vez de marcado.
    Defensas (hallazgos de careo, 05/10): solo si el rango empieza en la
    propia ancla (como garantiza Excel), recortado al techo de seguridad,
    con un presupuesto de celdas por hoja (muchas anclas hostiles
    multiplicaban el trabajo), sin celdas combinadas (MergedCell no admite
    la marca) y sin crear celdas al consultar."""
    try:
        min_c, min_f, max_c, max_f = range_boundaries(ref)
    except (ValueError, TypeError):
        return presupuesto, 0, 0
    if (min_f, min_c) != (ancla.row, ancla.column):
        return presupuesto, 0, 0
    max_f, max_c = min(max_f, LIMITE_FILAS), min(max_c, LIMITE_COLUMNAS)
    if (max_f - min_f + 1) * (max_c - min_c + 1) > presupuesto:
        # No cabe: solo la ancla, y convertir() lo avisa en el .md -- ni
        # corte silencioso a medias ni miles de marcas (careo, BAJA, 05/10).
        grandes.append(ref)
        return presupuesto, 0, 0
    f_max = c_max = 0
    for f in range(min_f, max_f + 1):
        for col in range(min_c, max_c + 1):
            if presupuesto <= 0:
                return presupuesto, f_max, c_max
            presupuesto -= 1  # cuenta lo recorrido, no solo lo marcado
            if (f, col) in sin_valor:
                continue
            celda = hoja._cells.get((f, col))
            if isinstance(celda, MergedCell) or (celda is not None and celda.value is not None):
                continue
            sin_valor[(f, col)] = None
            f_max, c_max = max(f_max, f), max(c_max, col)
    return presupuesto, f_max, c_max


def convertir(origen: Path) -> tuple[str, int]:
    """Devuelve (markdown, nº de celdas con fórmula sin valor cacheado).
    El segundo valor es un aviso -- 0 no garantiza que no falte nada, solo
    que no se detectó ninguna fórmula sin recalcular.
    Un .xls (Excel 97-2003) se pasa antes a .xlsx con LibreOffice: openpyxl
    no lee el formato binario, y así se conservan las combinadas."""
    if origen.suffix.lower() == ".xls":
        with tempfile.TemporaryDirectory(prefix="all2md-") as tmp:
            return _convertir_xlsx(con_soffice(origen, "xlsx", Path(tmp)), origen)
    return _convertir_xlsx(origen, origen)


def _convertir_xlsx(origen: Path, original: Path) -> tuple[str, int]:
    wb = load_workbook(str(origen), data_only=True)
    wb_formulas = load_workbook(str(origen), data_only=False)

    # Primera pasada: rango útil de cada hoja y estado de sus fórmulas.
    hojas = []
    hojas_omitidas = []
    matriciales_grandes = []
    total_sin_valor = total_con_valor = 0
    for hoja, hoja_formulas in zip(wb.worksheets, wb_formulas.worksheets):
        if hoja.max_row > LIMITE_FILAS or hoja.max_column > LIMITE_COLUMNAS:
            hojas_omitidas.append(
                f"{hoja.title} ({hoja.max_row}x{hoja.max_column}, supera "
                f"el límite de seguridad {LIMITE_FILAS}x{LIMITE_COLUMNAS})"
            )
            continue
        max_fila, max_col, sin_valor, con_valor, grandes = _escanear(hoja, hoja_formulas)
        matriciales_grandes += [f"{hoja.title}!{ref}" for ref in grandes]
        # Una cabecera combinada con contenido cubre todo su rango, aunque
        # sus celdas interiores (MergedCell) estén vacías: sin esto se
        # recortaba si iba más allá de la última celda con dato.
        for rango in hoja.merged_cells.ranges:
            ancla = hoja_formulas.cell(row=rango.min_row, column=rango.min_col).value
            if ancla is not None:
                max_fila = max(max_fila, rango.max_row)
                max_col = max(max_col, rango.max_col)
        total_sin_valor += len(sin_valor)
        total_con_valor += con_valor
        hojas.append((hoja, max_fila, max_col, sin_valor))

    # Una fórmula que devuelve "" (=SI(...;"";x), habitual para ocultar
    # ceros) llega como None igual que una sin calcular. Solo se puede
    # distinguir a nivel de libro: si ALGUNA fórmula trae valor, el libro se
    # recalculó y esos None son "" legítimos (hallazgo de careo, MEDIA,
    # 04/10). Excel recalcula el libro entero al guardar.
    libro_sin_recalcular = total_sin_valor > 0 and total_con_valor == 0

    partes = []
    for hoja, max_fila, max_col, sin_valor in hojas:
        if max_fila == 0 or max_col == 0:
            continue  # hoja vacía

        if libro_sin_recalcular:
            # Antes de propagar combinadas: así una cabecera combinada que
            # sea fórmula repite la marca en todas sus celdas.
            for fila, col in sin_valor:
                hoja.cell(row=fila, column=col).value = MARCA_FORMULA

        # Propaga el valor de cada rango combinado a todas sus celdas --
        # esto es lo que elimina las columnas "Unnamed": la cabecera ya no
        # está vacía porque Excel la pintaba combinada, ahora está repetida
        # en cada celda.
        for rango in list(hoja.merged_cells.ranges):
            ancla = hoja.cell(row=rango.min_row, column=rango.min_col)
            valor, formato = ancla.value, ancla.number_format
            hoja.unmerge_cells(str(rango))
            for fila in range(rango.min_row, rango.max_row + 1):
                for col in range(rango.min_col, rango.max_col + 1):
                    celda = hoja.cell(row=fila, column=col)
                    celda.value = valor
                    # El formato viaja con el valor: un 25% combinado no
                    # puede salir como 0.25 en sus celdas interiores.
                    celda.number_format = formato

        rejilla = [
            [_formatea(c.value, c.number_format) for c in fila]
            for fila in hoja.iter_rows(min_row=1, max_row=max_fila, min_col=1, max_col=max_col)
        ]
        # Filas y columnas separadoras (habituales en un USALI): "|  |" no
        # aporta nada a una IA/RAG y solo gasta tokens. Se decide DESPUÉS de
        # propagar combinadas: una columna cubierta por una cabecera
        # combinada ya no está vacía.
        rejilla = [fila for fila in rejilla if any(fila)]
        if not rejilla:
            # Solo tenía celdas de puro espacio: sin esto quedaba un
            # "## Título" huérfano y el libro no contaba como vacío.
            continue
        columnas = [i for i in range(max_col) if any(fila[i] for fila in rejilla)]

        # Una hoja oculta suele ser una tabla auxiliar de cálculo: se
        # incluye (puede tener datos), pero señalada para que la IA/RAG no
        # la tome por parte del informe que ve el lector.
        oculta = " (hoja oculta)" if hoja.sheet_state != "visible" else ""
        partes.append(f"## {hoja.title}{oculta}\n")
        for fila in rejilla:
            partes.append("| " + " | ".join(fila[i] for i in columnas) + " |")
        partes.append("")

    if libro_sin_recalcular and matriciales_grandes:
        partes.append(
            "**AVISO: fórmula(s) matricial(es) sin calcular demasiado grandes "
            "para marcar celda a celda -- solo se marca su primera celda:**\n- "
            + "\n- ".join(matriciales_grandes)
        )

    if hojas_omitidas:
        partes.append(
            "**AVISO: hoja(s) omitidas por superar el límite de seguridad -- "
            "revisar a mano si contienen algo relevante:**\n- "
            + "\n- ".join(hojas_omitidas)
        )

    md = "\n".join(partes)
    if not md.strip():
        raise ConversionVacia(f"{original.name} no produjo ningún contenido (¿libro vacío?).")
    return md, total_sin_valor if libro_sin_recalcular else 0


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit("Uso: python motor_excel.py <origen.xlsx> <destino.md>")
    origen = Path(sys.argv[1]).expanduser()
    destino = Path(sys.argv[2]).expanduser()
    if not origen.exists():
        sys.exit(f"No existe: {origen}")
    if misma_ruta(origen, destino):
        sys.exit(f"El destino es el propio documento de origen ({origen.name}): "
                 "no se sobrescribe.")

    try:
        md, avisos = convertir(origen)
    except ConversionVacia as e:
        sys.exit(str(e))
    except Exception as e:
        sys.exit(f"motor_excel falló con {origen.name}: {e}")

    try:
        escribir_atomico(destino, md)
    except (OSError, UnicodeError) as e:
        # Sin esto, un destino no escribible acababa en traceback crudo,
        # que la app enseña tal cual en el diálogo de error.
        sys.exit(f"No se pudo escribir {destino}: {e}")
    print(f"Guardado en {destino} ({len(md)} caracteres)")
    if avisos:
        print(
            f"AVISO: {avisos} celda(s) con fórmula sin valor cacheado "
            f"(marcadas como {MARCA_FORMULA!r} -- abrir y guardar el libro "
            "en Excel para recalcularlas)."
        )
